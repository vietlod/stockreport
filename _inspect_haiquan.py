"""Focused analysis of haiquan.xlsx + customs.gov.vn API testing."""
import openpyxl
import json
import re
from collections import Counter
from urllib.parse import urlparse, unquote

wb = openpyxl.load_workbook('docs/haiquan/haiquan.xlsx')
ws = wb.active

# Collect all URLs
all_sb = []
all_ct = []
for row_idx in range(2, ws.max_row + 1):
    sb = ws.cell(row=row_idx, column=4).value
    ct = ws.cell(row=row_idx, column=6).value
    thongtin = ws.cell(row=row_idx, column=1).value or ''
    thang = ws.cell(row=row_idx, column=2).value or ''
    if sb and str(sb).strip().startswith('http'):
        all_sb.append({'url': str(sb).strip(), 'thongtin': str(thongtin), 'thang': str(thang)})
    if ct and str(ct).strip().startswith('http'):
        all_ct.append({'url': str(ct).strip(), 'thongtin': str(thongtin), 'thang': str(thang)})

# === Classify URLs ===
def classify(url):
    if '/CustomsCMS/' in url:
        return 'CustomsCMS'
    elif '/Lists/ThongKeHaiQuanLichCongBo/Attachments/' in url:
        return 'SP-Lists'
    elif '/DocLib/' in url:
        return 'SP-DocLib'
    return 'Other'

sb_types = Counter(classify(i['url']) for i in all_sb)
ct_types = Counter(classify(i['url']) for i in all_ct)
print("=== URL CLASSIFICATION ===")
print(f"LinkSB ({len(all_sb)} total): {dict(sb_types)}")
print(f"LinkCT ({len(all_ct)} total): {dict(ct_types)}")

# === Filename pattern with proper year extraction ===
print("\n=== YEAR DISTRIBUTION (from Thang column) ===")
year_dist = Counter()
for row_idx in range(2, ws.max_row + 1):
    thang = ws.cell(row=row_idx, column=2).value
    if thang:
        s = str(thang)
        m = re.search(r'(20\d{2})', s)
        if m:
            year_dist[m.group(1)] += 1
for y, c in sorted(year_dist.items()):
    print(f"  {y}: {c} rows")

# === Period & Report Type from filename ===
print("\n=== FILENAME CONVENTIONS ===")
# Pattern: {YYYY}-T{MM}{Period}-{ReportNum}{Type}({Lang}-{Status}).pdf
period_cnt = Counter()
report_cnt = Counter()
for item in all_sb + all_ct:
    fn = unquote(item['url'].split('/')[-1])
    m = re.search(r'(\d{4})-T(\d{2})(T|K\d)-(\d+)([A-Z]+)\(VN-(SB|CT)\)', fn)
    if m:
        period_cnt[m.group(3)] += 1
        report_cnt[m.group(4) + m.group(5)] += 1

print("Period types (T=Tháng, K1=Nửa đầu tháng, K2=Nửa cuối tháng):")
for p, c in period_cnt.most_common():
    print(f"  {p}: {c}")

print("\nReport codes:")
for r, c in report_cnt.most_common(15):
    print(f"  {r}: {c}")

# === Example full URLs by category ===
print("\n=== EXAMPLE URLs ===")
for cat, label in [('CustomsCMS', 'New CMS'), ('SP-Lists', 'SharePoint Lists'), ('SP-DocLib', 'SharePoint DocLib')]:
    sb_ex = [i for i in all_sb if classify(i['url']) == cat][:2]
    ct_ex = [i for i in all_ct if classify(i['url']) == cat][:2]
    if sb_ex or ct_ex:
        print(f"\n{label}:")
        for i in sb_ex:
            print(f"  SB: {i['url']}")
        for i in ct_ex:
            print(f"  CT: {i['url']}")

# === Data quality summary ===
print("\n=== DATA QUALITY ===")
empty_info = sum(1 for r in range(2, ws.max_row+1) if not (ws.cell(row=r, column=1).value or '').strip())
both = sum(1 for r in range(2, ws.max_row+1)
    if (ws.cell(row=r, column=4).value or '').strip().startswith('http')
    and (ws.cell(row=r, column=6).value or '').strip().startswith('http'))
only_sb = sum(1 for r in range(2, ws.max_row+1)
    if (ws.cell(row=r, column=4).value or '').strip().startswith('http')
    and not (ws.cell(row=r, column=6).value or '').strip().startswith('http'))
only_ct = sum(1 for r in range(2, ws.max_row+1)
    if not (ws.cell(row=r, column=4).value or '').strip().startswith('http')
    and (ws.cell(row=r, column=6).value or '').strip().startswith('http'))

print(f"Total rows: {ws.max_row - 1}")
print(f"Rows with empty Thongtin: {empty_info}")
print(f"Both SB+CT: {both}, Only SB: {only_sb}, Only CT: {only_ct}")

# === Test URL accessibility ===
print("\n=== URL ACCESSIBILITY TEST ===")
import urllib.request
test_urls = [
    ("PDF (new CMS)", "https://files.customs.gov.vn/CustomsCMS/TONG_CUC/2022/9/20/2022-T09K1-1X(VN-SB).pdf"),
    ("PDF (SP Lists)", all_ct[-10]['url'] if len(all_ct) > 10 else ""),
]
for label, url in test_urls:
    if not url:
        continue
    try:
        req = urllib.request.Request(url, method='HEAD')
        req.add_header('User-Agent', 'Mozilla/5.0')
        resp = urllib.request.urlopen(req, timeout=10)
        print(f"  {label}: {resp.status} | Content-Type: {resp.headers.get('Content-Type', 'N/A')} | Size: {resp.headers.get('Content-Length', 'N/A')}")
    except Exception as e:
        print(f"  {label}: FAILED - {type(e).__name__}: {e}")
    print(f"    URL: {url}")

print("\n=== DONE ===")
