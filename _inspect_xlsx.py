"""Quick inspection of stock data Excel files."""
import openpyxl

files = [
    "docs/stock_exchange.xlsx",
    "docs/stock_industry.xlsx",
    "docs/stock_index.xlsx",
]

for f in files:
    print(f"\n{'='*60}")
    print(f"FILE: {f}")
    print('='*60)
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"  Sheet: {sname} | Rows: {ws.max_row} | Cols: {ws.max_column}")
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if i >= 7:
                break
        for r in rows:
            print(f"    {r}")
    wb.close()
