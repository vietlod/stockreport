"""Merge haiquan.xlsx + haiquan2226.xlsx → haiquan0925.xlsx

Đọc cả hai file xlsx, extract tất cả PDF URLs,
deduplicate theo URL, ghi ra file tổng hợp với schema thống nhất.

Schema output: [URL, Source]
- Mỗi row = 1 unique URL
- Source = tên file gốc (haiquan.xlsx hoặc haiquan2226.xlsx)

Chạy: python _merge_xlsx.py
"""
import sys
import openpyxl
from pathlib import Path

DOCS_DIR = Path(__file__).parent / "docs" / "haiquan"
FILE_1 = DOCS_DIR / "haiquan.xlsx"
FILE_2 = DOCS_DIR / "haiquan2226.xlsx"
OUTPUT = DOCS_DIR / "haiquan0925.xlsx"


def extract_urls_from_xlsx(xlsx_path: Path) -> list[dict]:
    """Extract tất cả HTTP URLs từ xlsx file.
    
    Tìm columns có header LinkSB/LinkCT, nếu không có thì quét tất cả cells.
    """
    print(f"  📖 Reading {xlsx_path.name}...")
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    ws = wb.active

    # Get headers
    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(c).strip() if c else '' for c in first_row]
    print(f"     Headers: {headers}")

    # Find URL columns (LinkSB, LinkCT) or auto-detect
    url_cols = []
    for i, h in enumerate(headers):
        if h in ("LinkSB", "LinkCT", "Sơ bộ", "Điều chỉnh", "Chính thức"):
            url_cols.append(i)
    
    auto_detect = len(url_cols) == 0
    if auto_detect:
        print(f"     ⚠ No known URL columns → auto-detecting from all cells")

    results = []
    seen = set()
    row_count = 0
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_count += 1
        cols_to_check = range(len(row)) if auto_detect else url_cols
        for idx in cols_to_check:
            if idx >= len(row):
                continue
            val = row[idx]
            if not isinstance(val, str) or not val.startswith("http"):
                continue
            url = val.strip()
            if url in seen or url == "null":
                continue
            seen.add(url)
            results.append({
                "url": url,
                "source": xlsx_path.name,
            })

    wb.close()
    print(f"     ✔ {len(results)} unique URLs from {row_count} rows")
    return results


def main():
    all_urls = []
    seen_urls = set()

    for f in [FILE_1, FILE_2]:
        if not f.exists():
            print(f"  ⚠ File not found: {f}")
            continue
        urls = extract_urls_from_xlsx(f)
        for u in urls:
            if u["url"] not in seen_urls:
                all_urls.append(u)
                seen_urls.add(u["url"])

    print(f"\n📊 Total unique URLs: {len(all_urls)}")

    # Write output
    print(f"📝 Writing {OUTPUT.name}...")
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "HaiQuan_0925"
    ws_out.append(["URL", "Source"])

    for u in all_urls:
        ws_out.append([u["url"], u["source"]])

    wb_out.save(str(OUTPUT))
    wb_out.close()
    print(f"✅ Done! {OUTPUT} ({len(all_urls)} URLs)")


if __name__ == "__main__":
    main()
